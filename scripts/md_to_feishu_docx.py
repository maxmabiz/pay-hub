#!/usr/bin/env python3
"""Convert the payment-center PRD markdown into a Feishu-importable .docx."""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "docs" / "付款中心需求说明书.md"
OUTS = [
    ROOT / "docs" / "付款中心需求说明书.docx",
    ROOT / "docs" / "feishu" / "付款中心需求说明书.docx",
]


def set_run_font(run, name="微软雅黑", size=11, bold=False, color=None):
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = name
    rpr = run._element.rPr
    if rpr is None:
        rpr = OxmlElement("w:rPr")
        run._element.insert(0, rpr)
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:ascii"), name)
    rfonts.set(qn("w:hAnsi"), name)
    rfonts.set(qn("w:eastAsia"), name)
    if color:
        run.font.color.rgb = RGBColor(*color)


def add_paragraph(doc, text, style="Normal", *, size=11, bold=False, color=None, space_after=8):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(0)
    run = p.add_run(text)
    heading_size = {"Title": 22, "Heading 1": 18, "Heading 2": 14, "Heading 3": 12, "Heading 4": 11}
    set_run_font(run, size=heading_size.get(style, size), bold=bold or style.startswith("Heading") or style == "Title", color=color)
    return p


def add_table(doc, rows):
    if not rows:
        return
    cols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=cols)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, row in enumerate(rows):
        for j in range(cols):
            cell = table.cell(i, j)
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            run = p.add_run(row[j] if j < len(row) else "")
            set_run_font(run, size=9, bold=(i == 0))
            if i == 0:
                tc_pr = cell._tc.get_or_add_tcPr()
                shd = OxmlElement("w:shd")
                shd.set(qn("w:fill"), "1F4E79")
                shd.set(qn("w:val"), "clear")
                tc_pr.append(shd)
                run.font.color.rgb = RGBColor(255, 255, 255)
    doc.add_paragraph().paragraph_format.space_after = Pt(6)


def parse_table(lines, start):
    rows = []
    i = start
    while i < len(lines) and lines[i].startswith("|"):
        raw = [c.strip() for c in lines[i].strip().strip("|").split("|")]
        if not re.match(r"^:?-{3,}:?$", raw[0] or ""):
            rows.append(raw)
        i += 1
    return rows, i


def convert(md_text: str, dest: Path) -> None:
    dest.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)

    style = doc.styles["Normal"]
    style.font.name = "微软雅黑"
    style.font.size = Pt(11)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")

    lines = md_text.replace("\r\n", "\n").split("\n")
    i = 0
    first_h1 = True
    in_code = False
    code_buf: list[str] = []

    while i < len(lines):
        line = lines[i]
        if line.startswith("```"):
            if in_code:
                add_paragraph(doc, "\n".join(code_buf), size=10, space_after=10)
                code_buf = []
                in_code = False
            else:
                in_code = True
            i += 1
            continue
        if in_code:
            code_buf.append(line)
            i += 1
            continue
        if not line.strip() or line.strip() == "---":
            i += 1
            continue
        if line.startswith("|"):
            rows, i = parse_table(lines, i)
            add_table(doc, rows)
            continue
        if line.startswith("# "):
            add_paragraph(doc, line[2:].strip(), "Title" if first_h1 else "Heading 1", color=(31, 78, 121))
            first_h1 = False
            i += 1
            continue
        if line.startswith("#### "):
            add_paragraph(doc, line[5:].strip(), "Heading 4", color=(31, 78, 121))
            i += 1
            continue
        if line.startswith("### "):
            add_paragraph(doc, line[4:].strip(), "Heading 3", color=(31, 78, 121))
            i += 1
            continue
        if line.startswith("## "):
            add_paragraph(doc, line[3:].strip(), "Heading 2", color=(31, 78, 121))
            i += 1
            continue
        if line.startswith("- "):
            add_paragraph(doc, "• " + line[2:].strip(), size=11)
            i += 1
            continue
        if re.match(r"^\d+\. ", line):
            add_paragraph(doc, line.strip(), size=11)
            i += 1
            continue
        add_paragraph(doc, line.strip(), size=11)
        i += 1

    doc.save(dest)


def main():
    text = SRC.read_text(encoding="utf-8")
    for dest in OUTS:
        convert(text, dest)
        print(f"wrote {dest}")


if __name__ == "__main__":
    main()
